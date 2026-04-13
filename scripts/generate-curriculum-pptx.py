#!/usr/bin/env python3
"""Generate SCAiL-branded PowerPoint decks for every SCAiL course.

Reads slide content from:
  scripts/data/curriculum-slides-batch1.json   (sc-1 through sc-4)
  scripts/data/curriculum-slides-batch2.json   (sc-5 through sc-8)

Writes .pptx files to public/docs/curriculum/ so the tracker can serve
them as downloads. Also writes public/docs/curriculum/index.json listing
every deck with filename, title, slide count, and file size.
"""

import json
import os
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# ---------- Brand ----------
NAVY = RGBColor(0x1B, 0x36, 0x5D)
GOLD = RGBColor(0xD4, 0xA8, 0x43)
TEAL = RGBColor(0x2E, 0xC4, 0xB6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY = RGBColor(0xF8, 0xFA, 0xFC)
SLATE = RGBColor(0x33, 0x41, 0x55)
NEAR_BLACK = RGBColor(0x0F, 0x17, 0x2A)

FONT = "Inter"

# 16:9
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = ROOT / "scripts" / "data"
OUT_DIR = ROOT / "public" / "docs" / "curriculum"
OUT_DIR.mkdir(parents=True, exist_ok=True)


# ---------- Helpers ----------
def set_bg(slide, color):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = color
    bg.line.fill.background()
    bg.shadow.inherit = False
    return bg


def add_rect(slide, x, y, w, h, color, line=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    if not line:
        shape.line.fill.background()
    return shape


def add_text(slide, x, y, w, h, text, *, size=18, bold=False, color=SLATE,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font=FONT):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return tb


def add_bullets(slide, x, y, w, h, bullets, *, size=20, color=SLATE, spacing=8):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    for i, bullet in enumerate(bullets):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.space_after = Pt(spacing)
        p.alignment = PP_ALIGN.LEFT
        run = p.add_run()
        run.text = f"•  {bullet}"
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.color.rgb = color
    return tb


def add_footer(slide, course_title):
    # Gold accent line
    add_rect(slide, Inches(0.5), Inches(7.05), Inches(2), Emu(18000), GOLD)
    add_text(
        slide, Inches(0.5), Inches(7.1), Inches(8), Inches(0.3),
        f"SCAiL  ·  {course_title}",
        size=9, color=SLATE,
    )
    add_text(
        slide, Inches(11.5), Inches(7.1), Inches(1.5), Inches(0.3),
        "scail.org",
        size=9, color=SLATE, align=PP_ALIGN.RIGHT,
    )


def add_wordmark(slide, x, y, *, light=False):
    color_sc = WHITE if light else NAVY
    color_ai = TEAL
    tb = slide.shapes.add_textbox(x, y, Inches(2), Inches(0.5))
    tf = tb.text_frame
    tf.margin_left = Emu(0)
    tf.margin_top = Emu(0)
    p = tf.paragraphs[0]
    r1 = p.add_run()
    r1.text = "SC"
    r1.font.name = FONT
    r1.font.size = Pt(24)
    r1.font.bold = True
    r1.font.color.rgb = color_sc
    r2 = p.add_run()
    r2.text = "Ai"
    r2.font.name = FONT
    r2.font.size = Pt(24)
    r2.font.bold = True
    r2.font.color.rgb = color_ai
    r3 = p.add_run()
    r3.text = "L"
    r3.font.name = FONT
    r3.font.size = Pt(24)
    r3.font.bold = True
    r3.font.color.rgb = color_sc


# ---------- Slide renderers ----------
def render_title(prs, slide_data, course_title):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    set_bg(slide, NAVY)

    # Gold accent bar
    add_rect(slide, Inches(0.75), Inches(2.1), Inches(0.8), Inches(0.08), GOLD)

    add_text(
        slide, Inches(0.75), Inches(2.3), Inches(12), Inches(1.6),
        slide_data.get("title", ""),
        size=60, bold=True, color=WHITE,
    )
    if slide_data.get("subtitle"):
        add_text(
            slide, Inches(0.75), Inches(4.2), Inches(12), Inches(1),
            slide_data["subtitle"],
            size=28, color=GOLD,
        )
    footer_text = slide_data.get("footer", "SCAiL — South Carolina AI Literacy")
    add_text(
        slide, Inches(0.75), Inches(6.8), Inches(12), Inches(0.4),
        footer_text,
        size=12, color=WHITE,
    )
    # Teal corner accent
    add_rect(slide, Inches(12.2), Inches(0.5), Inches(0.6), Inches(0.6), TEAL)


def render_agenda(prs, slide_data, course_title):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, LIGHT_GRAY)
    add_wordmark(slide, Inches(0.5), Inches(0.4))

    add_rect(slide, Inches(0.5), Inches(1.6), Inches(1.2), Emu(30000), TEAL)
    add_text(
        slide, Inches(0.5), Inches(1.7), Inches(12), Inches(1),
        slide_data.get("title", "What We'll Cover Today"),
        size=40, bold=True, color=NAVY,
    )

    items = slide_data.get("items", [])
    if items:
        numbered = [f"{i+1}.  {item}" for i, item in enumerate(items)]
        add_bullets(
            slide, Inches(0.75), Inches(3.2), Inches(12), Inches(3.5),
            numbered, size=22, color=SLATE, spacing=14,
        )
    add_footer(slide, course_title)


def render_content(prs, slide_data, course_title):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, WHITE)
    add_wordmark(slide, Inches(0.5), Inches(0.4))

    add_rect(slide, Inches(0.5), Inches(1.5), Inches(0.8), Emu(25000), GOLD)
    add_text(
        slide, Inches(0.5), Inches(1.6), Inches(12), Inches(1),
        slide_data.get("title", ""),
        size=32, bold=True, color=NAVY,
    )

    bullets = slide_data.get("bullets", [])
    if bullets:
        add_bullets(
            slide, Inches(0.75), Inches(2.9), Inches(12), Inches(3.8),
            bullets, size=20, color=SLATE, spacing=10,
        )

    note = slide_data.get("note")
    if note:
        add_rect(slide, Inches(0.5), Inches(6.3), Inches(12.3), Inches(0.55), LIGHT_GRAY)
        add_rect(slide, Inches(0.5), Inches(6.3), Emu(40000), Inches(0.55), TEAL)
        add_text(
            slide, Inches(0.75), Inches(6.4), Inches(12), Inches(0.4),
            f"Presenter note: {note}",
            size=11, color=SLATE,
        )
    add_footer(slide, course_title)


def render_quote(prs, slide_data, course_title):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, NAVY)
    add_wordmark(slide, Inches(0.5), Inches(0.4), light=True)

    add_text(
        slide, Inches(0.75), Inches(1.5), Inches(2), Inches(1),
        "\u201C",
        size=90, bold=True, color=GOLD,
    )
    add_text(
        slide, Inches(1.5), Inches(2.7), Inches(10.5), Inches(3),
        slide_data.get("text", ""),
        size=28, color=WHITE, anchor=MSO_ANCHOR.TOP,
    )
    if slide_data.get("attribution"):
        add_text(
            slide, Inches(1.5), Inches(5.8), Inches(10.5), Inches(0.5),
            f"— {slide_data['attribution']}",
            size=14, color=GOLD,
        )
    add_footer(slide, course_title)


def render_activity(prs, slide_data, course_title):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, LIGHT_GRAY)
    add_wordmark(slide, Inches(0.5), Inches(0.4))

    # Teal header badge
    add_rect(slide, Inches(0.5), Inches(1.5), Inches(2.2), Inches(0.45), TEAL)
    add_text(
        slide, Inches(0.65), Inches(1.55), Inches(2), Inches(0.4),
        "HANDS-ON",
        size=11, bold=True, color=WHITE,
    )
    add_text(
        slide, Inches(0.5), Inches(2.1), Inches(12.5), Inches(1),
        slide_data.get("title", "Hands-On Activity"),
        size=32, bold=True, color=NAVY,
    )
    steps = slide_data.get("steps", [])
    if steps:
        numbered = [f"Step {i+1}.  {s}" for i, s in enumerate(steps)]
        add_bullets(
            slide, Inches(0.75), Inches(3.3), Inches(12), Inches(3.6),
            numbered, size=20, color=SLATE, spacing=12,
        )
    add_footer(slide, course_title)


def render_section(prs, slide_data, course_title):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, TEAL)
    add_wordmark(slide, Inches(0.5), Inches(0.4), light=True)

    add_text(
        slide, Inches(0.75), Inches(3.1), Inches(12), Inches(1.5),
        slide_data.get("title", ""),
        size=54, bold=True, color=WHITE,
    )
    add_rect(slide, Inches(0.75), Inches(4.8), Inches(1.2), Emu(35000), GOLD)


def render_closing(prs, slide_data, course_title):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, WHITE)
    add_wordmark(slide, Inches(0.5), Inches(0.4))

    add_rect(slide, Inches(0.5), Inches(1.5), Inches(0.8), Emu(25000), GOLD)
    add_text(
        slide, Inches(0.5), Inches(1.6), Inches(12), Inches(1),
        slide_data.get("title", "What You Walk Away With"),
        size=36, bold=True, color=NAVY,
    )
    bullets = slide_data.get("bullets", [])
    if bullets:
        checked = [f"\u2713  {b}" for b in bullets]
        add_bullets(
            slide, Inches(0.75), Inches(2.9), Inches(12), Inches(3.5),
            checked, size=22, color=SLATE, spacing=12,
        )

    # CTA band
    add_rect(slide, 0, Inches(6.2), SLIDE_W, Inches(0.9), NAVY)
    add_text(
        slide, Inches(0.75), Inches(6.35), Inches(12), Inches(0.6),
        "Questions? hello@scail.org  ·  scail.org",
        size=16, color=WHITE,
    )
    add_text(
        slide, Inches(0.75), Inches(6.7), Inches(12), Inches(0.4),
        "AI Literacy for Every Community",
        size=11, color=GOLD,
    )


RENDERERS = {
    "title": render_title,
    "agenda": render_agenda,
    "content": render_content,
    "quote": render_quote,
    "activity": render_activity,
    "section": render_section,
    "closing": render_closing,
}


# ---------- Main ----------
def load_slides():
    merged = {}
    for batch in ("curriculum-slides-batch1.json", "curriculum-slides-batch2.json"):
        path = DATA_DIR / batch
        if not path.exists():
            print(f"  [warn] missing {path}")
            continue
        with open(path, "r") as f:
            data = json.load(f)
        merged.update(data)
    return merged


COURSE_META = {
    "sc-1": {"short": "AI-101", "filename": "SCAiL-AI-101.pptx"},
    "sc-2": {"short": "AI-for-Your-Job", "filename": "SCAiL-AI-for-Your-Job.pptx"},
    "sc-3": {"short": "AI-for-Small-Business", "filename": "SCAiL-AI-for-Small-Business.pptx"},
    "sc-4": {"short": "AI-for-Farmers", "filename": "SCAiL-AI-for-Farmers.pptx"},
    "sc-5": {"short": "AI-Safety", "filename": "SCAiL-AI-Safety.pptx"},
    "sc-6": {"short": "Youth-AI-Lab", "filename": "SCAiL-Youth-AI-Lab.pptx"},
    "sc-7": {"short": "AI-Leadership-Academy", "filename": "SCAiL-AI-Leadership-Academy.pptx"},
    "sc-8": {"short": "Train-the-Trainer", "filename": "SCAiL-Train-the-Trainer.pptx"},
}


def build_deck(course_id, course_data):
    course_title = course_data.get("title", course_id)
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    slides = course_data.get("slides", [])
    for s in slides:
        renderer = RENDERERS.get(s.get("type"))
        if renderer is None:
            continue
        renderer(prs, s, course_title)

    meta = COURSE_META.get(course_id, {"filename": f"{course_id}.pptx"})
    out = OUT_DIR / meta["filename"]
    prs.save(out)
    return out, len(slides)


def main():
    slides_data = load_slides()
    if not slides_data:
        print("No slide data found. Agents haven't finished yet?")
        return

    index = []
    for course_id in sorted(slides_data.keys()):
        course_data = slides_data[course_id]
        out, slide_count = build_deck(course_id, course_data)
        size_kb = round(os.path.getsize(out) / 1024, 1)
        index.append({
            "id": course_id,
            "title": course_data.get("title", course_id),
            "filename": out.name,
            "slides": slide_count,
            "sizeKB": size_kb,
        })
        print(f"  [ok] {course_id}  ->  {out.name}  ({slide_count} slides, {size_kb} KB)")

    with open(OUT_DIR / "index.json", "w") as f:
        json.dump(index, f, indent=2)
    print(f"\nWrote {len(index)} decks to {OUT_DIR}")
    print(f"Index: {OUT_DIR / 'index.json'}")


if __name__ == "__main__":
    main()
