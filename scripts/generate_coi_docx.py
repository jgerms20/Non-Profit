#!/usr/bin/env python3
"""Generate a Word document for SCAiL Conflict of Interest Policy."""

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

NAVY = RGBColor(0x1B, 0x36, 0x5D)
GOLD = RGBColor(0xD4, 0xA8, 0x43)
DARK = RGBColor(0x33, 0x33, 0x33)
GRAY = RGBColor(0x66, 0x66, 0x66)

style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)

def title(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(22)
    r.font.color.rgb = NAVY

def subtitle(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.italic = True
    r.font.size = Pt(12)
    r.font.color.rgb = GRAY

def article_h(text):
    h = doc.add_heading(text, level=1)
    for r in h.runs:
        r.font.color.rgb = NAVY

def section_h(text):
    h = doc.add_heading(text, level=2)
    for r in h.runs:
        r.font.color.rgb = GOLD
        r.font.size = Pt(12)

def body(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.color.rgb = DARK
    r.font.size = Pt(11)

def body_indent(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.5)
    r = p.add_run(text)
    r.font.color.rgb = DARK
    r.font.size = Pt(11)

def bullet(text):
    p = doc.add_paragraph(style='List Bullet')
    r = p.add_run(text)
    r.font.size = Pt(11)
    r.font.color.rgb = DARK

# ========== CONTENT ==========

title("CONFLICT OF INTEREST POLICY")
title("THE SCAIL INITIATIVE, INC.")
subtitle("A South Carolina Nonprofit Corporation")
doc.add_paragraph()
subtitle("Adopted: _____________, 2026")
subtitle("SC Filing ID: 260320-1607524")

doc.add_page_break()

# ARTICLE I
article_h("Article I — Purpose")
body("The purpose of this Conflict of Interest Policy is to protect the interests of The SCAiL Initiative, Inc. (the \"Corporation\") when it is contemplating entering into a transaction or arrangement that might benefit the private interest of a director, officer, or key employee of the Corporation, or might result in a possible excess benefit transaction.")
body("This policy is intended to supplement, but not replace, any applicable state and federal laws governing conflict of interest applicable to nonprofit and charitable organizations.")

# ARTICLE II
article_h("Article II — Definitions")

section_h("2.1 — Interested Person")
body("Any director, officer, or member of a committee with board-delegated powers who has a direct or indirect financial interest, as defined below, is an interested person.")

section_h("2.2 — Financial Interest")
body("A person has a financial interest if the person has, directly or indirectly, through business, investment, or family:")
bullet("An ownership or investment interest in any entity with which the Corporation has a transaction or arrangement;")
bullet("A compensation arrangement with the Corporation or with any entity or individual with which the Corporation has a transaction or arrangement; or")
bullet("A potential ownership or investment interest in, or compensation arrangement with, any entity or individual with which the Corporation is negotiating a transaction or arrangement.")
body("Compensation includes direct and indirect remuneration as well as gifts or favors that are not insubstantial. A financial interest is not necessarily a conflict of interest. A person who has a financial interest may have a conflict of interest only if the Board of Directors or appropriate committee decides that a conflict of interest exists, in accordance with this policy.")

# ARTICLE III
article_h("Article III — Procedures")

section_h("3.1 — Duty to Disclose")
body("In connection with any actual or possible conflict of interest, an interested person must disclose the existence of the financial interest and be given the opportunity to disclose all material facts to the Board of Directors or committee considering the proposed transaction or arrangement.")

section_h("3.2 — Determining Whether a Conflict of Interest Exists")
body("After disclosure of the financial interest and all material facts, and after any discussion with the interested person, the interested person shall leave the Board or committee meeting while the determination of a conflict of interest is discussed and voted upon. The remaining Board or committee members shall decide if a conflict of interest exists.")

section_h("3.3 — Procedures for Addressing the Conflict of Interest")
body("If a conflict of interest is determined to exist:")
bullet("An interested person may make a presentation at the Board or committee meeting, but after the presentation, the interested person shall leave the meeting during the discussion of, and the vote on, the transaction or arrangement involving the possible conflict of interest.")
bullet("The President or committee chair shall, if appropriate, appoint a disinterested person or committee to investigate alternatives to the proposed transaction or arrangement.")
bullet("After exercising due diligence, the Board or committee shall determine whether the Corporation can obtain with reasonable efforts a more advantageous transaction or arrangement from a person or entity that would not give rise to a conflict of interest.")
bullet("If a more advantageous transaction or arrangement is not reasonably possible under circumstances not producing a conflict of interest, the Board or committee shall determine by a majority vote of the disinterested directors whether the transaction or arrangement is in the Corporation's best interest, for its own benefit, and whether it is fair and reasonable. In conformity with the above determination, it shall make its decision as to whether to enter into the transaction or arrangement.")

section_h("3.4 — Violations of the Conflict of Interest Policy")
body("If the Board or committee has reasonable cause to believe a person has failed to disclose actual or possible conflicts of interest, it shall inform the person of the basis for such belief and afford the person an opportunity to explain the alleged failure to disclose.")
body("If, after hearing the person's response and after making further investigation as warranted by the circumstances, the Board or committee determines the person has failed to disclose an actual or possible conflict of interest, it shall take appropriate disciplinary and corrective action.")

# ARTICLE IV
article_h("Article IV — Records of Proceedings")
body("The minutes of the Board and all committees with board-delegated powers shall contain:")
bullet("The names of the persons who disclosed or otherwise were found to have a financial interest in connection with an actual or possible conflict of interest, the nature of the financial interest, any action taken to determine whether a conflict of interest was present, and the Board's or committee's decision as to whether a conflict of interest in fact existed.")
bullet("The names of the persons who were present for discussions and votes relating to the transaction or arrangement, the content of the discussion, including any alternatives to the proposed transaction or arrangement, and a record of any votes taken in connection with the proceedings.")

# ARTICLE V
article_h("Article V — Compensation")
body("A voting member of the Board who receives compensation, directly or indirectly, from the Corporation for services is precluded from voting on matters pertaining to that member's compensation.")
body("A voting member of any committee whose jurisdiction includes compensation matters and who receives compensation, directly or indirectly, from the Corporation for services is precluded from voting on matters pertaining to that member's compensation.")
body("No voting member of the Board or any committee whose jurisdiction includes compensation matters and who receives compensation, directly or indirectly, from the Corporation, either individually or collectively, is prohibited from providing information to any committee regarding compensation.")

# ARTICLE VI
article_h("Article VI — Annual Statements")
body("Each director, officer, and member of a committee with board-delegated powers shall annually sign a statement which affirms such person:")
bullet("Has received a copy of the Conflict of Interest Policy;")
bullet("Has read and understands the policy;")
bullet("Has agreed to comply with the policy; and")
bullet("Understands the Corporation is charitable and that in order to maintain its federal tax exemption it must engage primarily in activities which accomplish one or more of its tax-exempt purposes.")

# ARTICLE VII
article_h("Article VII — Periodic Reviews")
body("To ensure the Corporation operates in a manner consistent with charitable purposes and does not engage in activities that could jeopardize its tax-exempt status, periodic reviews shall be conducted. The periodic reviews shall, at a minimum, include the following subjects:")
bullet("Whether compensation arrangements and benefits are reasonable, based on competent survey information, and the result of arm's length bargaining.")
bullet("Whether partnerships, joint ventures, and arrangements with management organizations conform to the Corporation's written policies, are properly recorded, reflect reasonable investment or payments for goods and services, further charitable purposes, and do not result in inurement, impermissible private benefit, or an excess benefit transaction.")

doc.add_page_break()

# ANNUAL DISCLOSURE FORM
article_h("Annual Conflict of Interest Disclosure Form")
doc.add_paragraph()

body("Name: _______________________________________________")
body("Position: ____________________________________________")
body("Date: ________________________________________________")
doc.add_paragraph()

section_h("Part A — Acknowledgment")
body("I have received a copy of the Conflict of Interest Policy of The SCAiL Initiative, Inc.")
body("I have read and understand the policy.")
body("I agree to comply with the policy.")
body("I understand that The SCAiL Initiative, Inc. is a tax-exempt organization and that in order to maintain its federal tax exemption it must engage primarily in activities which accomplish one or more of its tax-exempt purposes.")
doc.add_paragraph()

section_h("Part B — Disclosure of Financial Interests")
body("Do you have any financial interests, as defined in Article II of the Conflict of Interest Policy, to disclose?")
doc.add_paragraph()
body("[ ] No, I have no financial interests to disclose.")
body("[ ] Yes, I have the following financial interests to disclose (describe below):")
doc.add_paragraph()
body("1. _____________________________________________________")
body("2. _____________________________________________________")
body("3. _____________________________________________________")
doc.add_paragraph()

section_h("Part C — Relationships")
body("Are you a director, officer, employee, or agent of any organization that does business with The SCAiL Initiative, Inc.?")
doc.add_paragraph()
body("[ ] No")
body("[ ] Yes (please describe): ________________________________")
doc.add_paragraph()
body("Do you have any family members who have a financial interest in any entity that does business with The SCAiL Initiative, Inc.?")
doc.add_paragraph()
body("[ ] No")
body("[ ] Yes (please describe): ________________________________")

doc.add_paragraph()
doc.add_paragraph()
body("Signature: ______________________________________________")
body("Printed Name: ___________________________________________")
body("Date: ___________________________________________________")

doc.add_paragraph()
doc.add_paragraph()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("The SCAiL Initiative, Inc. — South Carolina AI Literacy")
r.font.size = Pt(9)
r.font.color.rgb = GRAY
p2 = doc.add_paragraph()
p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r2 = p2.add_run("SC Filing ID: 260320-1607524")
r2.font.size = Pt(9)
r2.font.color.rgb = GRAY

doc.save('/home/user/Non-Profit/docs/SCAiL-Conflict-of-Interest-Policy.docx')
print("Done: docs/SCAiL-Conflict-of-Interest-Policy.docx")
