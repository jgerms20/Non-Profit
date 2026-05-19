#!/usr/bin/env python3
"""Generate a Word document for The SCAiL Initiative bylaws."""

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

NAVY = RGBColor(0x1B, 0x36, 0x5D)
GOLD = RGBColor(0xD4, 0xA8, 0x43)
DARK = RGBColor(0x33, 0x33, 0x33)
GRAY = RGBColor(0x66, 0x66, 0x66)

style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)

def title(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(24)
    r.font.color.rgb = NAVY

def subtitle(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.italic = True
    r.font.size = Pt(12)
    r.font.color.rgb = GRAY

def article(text):
    h = doc.add_heading(text, level=1)
    for r in h.runs:
        r.font.color.rgb = NAVY

def section(text):
    h = doc.add_heading(text, level=2)
    for r in h.runs:
        r.font.color.rgb = GOLD
        r.font.size = Pt(12)

def body(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.color.rgb = DARK
    r.font.size = Pt(11)
    return p

def body_indent(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.5)
    r = p.add_run(text)
    r.font.color.rgb = DARK
    r.font.size = Pt(11)

def sig_line(name, title_text):
    doc.add_paragraph()
    p = doc.add_paragraph()
    r = p.add_run(f"{name}")
    r.bold = True
    r.font.size = Pt(11)
    p2 = doc.add_paragraph()
    r2 = p2.add_run(title_text)
    r2.font.size = Pt(10)
    r2.font.color.rgb = GRAY
    p3 = doc.add_paragraph()
    r3 = p3.add_run("Date: _______________     Signature: _______________")
    r3.font.size = Pt(10)

# ========== CONTENT ==========

title("BYLAWS")
title("OF THE SCAIL INITIATIVE, INC.")
subtitle("A South Carolina Nonprofit Corporation")
doc.add_paragraph()
subtitle("Adopted: April 2026")
subtitle("SC Articles Filing ID: 260320-1607524")
subtitle("EIN: [INSERT — obtained May 2026]")

doc.add_page_break()

# ARTICLE I
article("ARTICLE I — NAME AND PRINCIPAL OFFICE")

section("Section 1.1 — Name")
body('The name of this corporation is The SCAiL Initiative, Inc. (hereinafter "the Corporation" or "SCAiL").')

section("Section 1.2 — Principal Office")
body("The principal office of the Corporation shall be located within the State of South Carolina, as designated by the Board of Directors. The Corporation may maintain additional offices or registered agents as the Board deems necessary.")

section("Section 1.3 — Registered Agent")
body("The Corporation shall maintain a registered agent in South Carolina as required by state law. The Board of Directors may change the registered agent by filing the required documents with the South Carolina Secretary of State.")

# ARTICLE II
article("ARTICLE II — MISSION AND PURPOSE")

section("Section 2.1 — Mission")
body("The SCAiL Initiative exists to advance AI literacy across South Carolina — particularly in underserved, rural, and historically overlooked communities — so that every person in the state has the knowledge and tools to participate in, benefit from, and shape the AI economy.")

section("Section 2.2 — Exempt Purpose")
body("The Corporation is organized exclusively for charitable, educational, and scientific purposes within the meaning of Section 501(c)(3) of the Internal Revenue Code (or the corresponding provision of any future United States revenue law), including but not limited to:")
body_indent("(a) Providing free and low-cost AI literacy education to community members across South Carolina;")
body_indent("(b) Designing and delivering curriculum for individuals without prior technical backgrounds;")
body_indent("(c) Training community educators, facilitators, and workforce professionals to teach AI literacy in their local contexts;")
body_indent("(d) Conducting research and collecting data on AI literacy gaps, community needs, and program effectiveness;")
body_indent("(e) Building partnerships with schools, libraries, community organizations, employers, and government agencies to expand access to AI education;")
body_indent("(f) Eventually supporting a community innovation incubator and investment network rooted in the communities the Corporation serves.")

section("Section 2.3 — Limitations")
body('No part of the net earnings of the Corporation shall inure to the benefit of, or be distributable to, its directors, officers, or other private persons, except that the Corporation shall be authorized and empowered to pay reasonable compensation for services rendered and to make payments and distributions in furtherance of the purposes set forth above. The Corporation shall not carry on any other activities not permitted to be carried on (a) by a corporation exempt from federal income tax under Section 501(c)(3) of the Internal Revenue Code, or (b) by a corporation contributions to which are deductible under Section 170(c)(2) of the Internal Revenue Code.')

# ARTICLE III
article("ARTICLE III — MEMBERSHIP")

section("Section 3.1 — No Members")
body("The Corporation shall have no members as that term is defined in the South Carolina Nonprofit Corporation Act. The Board of Directors shall have all rights and authority otherwise granted to members under the Act.")

# ARTICLE IV
article("ARTICLE IV — BOARD OF DIRECTORS")

section("Section 4.1 — General Powers")
body("The affairs of the Corporation shall be managed by its Board of Directors. The Board of Directors shall have control of and be responsible for the management of the affairs and funds of the Corporation.")

section("Section 4.2 — Number of Directors")
body("The Board of Directors shall consist of not fewer than three (3) and not more than fifteen (15) directors. The Board may, by resolution, increase or decrease the number of directors within these limits.")

section("Section 4.3 — Founding Board")
body("The founding Board of Directors consists of the following five (5) individuals:")

table = doc.add_table(rows=6, cols=3)
table.style = 'Light Grid Accent 1'
headers = ["Name", "Title", "Role"]
for i, h in enumerate(headers):
    cell = table.rows[0].cells[i]
    cell.text = ""
    r = cell.paragraphs[0].add_run(h)
    r.bold = True
    r.font.size = Pt(10)

board_data = [
    ("Joshua German", "Founder & President", "Executive Leadership"),
    ("Janel Moore", "Board Member, Treasurer", "Finance & Treasury"),
    ("Kinsey Meggett, PhD", "Board Member, Secretary", "Community Engagement & Education"),
    ("Tre Jenkins", "Board Member, Technology & Data", "Technology Strategy & Data"),
    ("Darren Burton", "Board Member, Education Policy & Operations", "Education Policy & Organizational Operations"),
]
for r_idx, (name, t, role) in enumerate(board_data, 1):
    for c_idx, val in enumerate([name, t, role]):
        cell = table.rows[r_idx].cells[c_idx]
        cell.text = ""
        run = cell.paragraphs[0].add_run(val)
        run.font.size = Pt(10)

section("Section 4.4 — Term of Office")
body("Directors shall serve two-year (2) terms. Directors may be re-elected without limitation. The initial board shall serve staggered terms as determined by the Board: Group A (Joshua German, Janel Moore, Kinsey Meggett) shall serve a three-year initial term; Group B (Tre Jenkins, Darren Burton) shall serve a two-year initial term. Thereafter, all terms shall be two years.")

section("Section 4.5 — Qualifications")
body("Directors shall be at least eighteen (18) years of age. Directors need not be residents of South Carolina.")

section("Section 4.6 — Election of Directors")
body("After the founding board, new directors shall be nominated by the Board and elected by a majority vote of the current Board of Directors.")

section("Section 4.7 — Vacancies")
body("Any vacancy occurring in the Board of Directors may be filled by the affirmative vote of a majority of the remaining directors. A director elected to fill a vacancy shall be elected for the unexpired portion of the term.")

section("Section 4.8 — Removal")
body("Any director may be removed from office, with or without cause, by the affirmative vote of two-thirds (2/3) of the remaining directors at any duly called meeting of the Board. A director who fails to attend three (3) consecutive meetings without excuse may be deemed to have resigned.")

section("Section 4.9 — Resignation")
body("Any director may resign at any time by giving written notice to the President or Secretary of the Corporation. Resignation is effective upon receipt unless otherwise specified.")

section("Section 4.10 — Compensation")
body("Directors shall serve without compensation for their service on the Board. Directors may be reimbursed for reasonable expenses incurred in the performance of their duties, as approved by the Board. This section shall not prevent directors from receiving reasonable compensation in another capacity (e.g., as a contractor or employee), provided such compensation is disclosed and approved in accordance with the Conflict of Interest Policy.")

# ARTICLE V
article("ARTICLE V — MEETINGS OF THE BOARD")

section("Section 5.1 — Annual Meeting")
body("The Board of Directors shall hold an annual meeting each calendar year at a time and place designated by the Board. The annual meeting shall include a review of the prior year's activities, financial reports, and election of directors as needed.")

section("Section 5.2 — Regular Meetings")
body("The Board shall meet no fewer than four (4) times per year (quarterly). Meeting dates for each calendar year shall be established at the prior year's annual meeting or by resolution of the Board.")

section("Section 5.3 — Special Meetings")
body("Special meetings of the Board of Directors may be called by the President or by any two (2) directors. Notice of a special meeting shall be given to each director at least forty-eight (48) hours in advance.")

section("Section 5.4 — Notice")
body("Notice of all Board meetings shall be given to each director by email, text message, or other electronic communication at least five (5) days before any regular or annual meeting. Notice shall include the date, time, location or video link, and agenda.")

section("Section 5.5 — Quorum")
body("A majority of the directors in office shall constitute a quorum for the transaction of business. If a quorum is not present, the meeting may be adjourned to a later date.")

section("Section 5.6 — Voting")
body("Each director shall have one (1) vote. Actions of the Board shall require an affirmative vote of a majority of directors present at a meeting at which a quorum is present, unless a greater vote is required by law or these Bylaws.")

section("Section 5.7 — Remote Participation")
body("Directors may participate in Board meetings by telephone, video conference, or other electronic means by which all participants can communicate with each other simultaneously. Participation by such means shall constitute presence at the meeting.")

section("Section 5.8 — Action Without Meeting")
body("The Board of Directors may take action without a meeting if all directors consent in writing (including by email) to the action. Written consent shall be filed with the minutes.")

# ARTICLE VI
article("ARTICLE VI — OFFICERS")

section("Section 6.1 — Officers")
body("The officers of the Corporation shall be a President, a Treasurer, and a Secretary. The Board may create additional officer positions as needed. No two offices may be held simultaneously by the same person.")

section("Section 6.2 — President")
body("The President shall be the chief executive officer of the Corporation, shall preside at all meetings of the Board of Directors, shall see that all orders and resolutions of the Board are carried out, and shall generally supervise the operations of the Corporation. The President shall execute contracts and other legal instruments on behalf of the Corporation when authorized by the Board.")

section("Section 6.3 — Treasurer")
body("The Treasurer shall have custody of all funds and securities of the Corporation, shall keep full and accurate accounts of receipts and disbursements, shall present financial reports at Board meetings, and shall perform such other duties as the Board may require. The Treasurer shall oversee grant reporting, budgeting, and financial compliance.")

section("Section 6.4 — Secretary")
body("The Secretary shall keep minutes of all meetings of the Board, maintain the records of the Corporation, give required notices, maintain the registered agent information, and perform such other duties as the Board may require.")

section("Section 6.5 — Election and Term")
body("Officers shall be elected by the Board of Directors and shall serve two-year (2) terms or until their successors are elected. Officers may serve consecutive terms.")

section("Section 6.6 — Removal")
body("Any officer may be removed by the Board of Directors with or without cause by a majority vote.")

section("Section 6.7 — Current Officers")
body("The founding officers are:")
body_indent("President: Joshua German")
body_indent("Treasurer: Janel Moore")
body_indent("Secretary: Kinsey Meggett, PhD")

# ARTICLE VII
article("ARTICLE VII — COMMITTEES")

section("Section 7.1 — Standing Committees")
body("The Board of Directors may establish standing committees to assist in carrying out the work of the Corporation. Current standing committees may include: Finance & Audit Committee, Program & Curriculum Committee, Technology & Data Committee, and Governance & Nominating Committee.")

section("Section 7.2 — Committee Composition")
body("Each committee shall have at least one Board member. Committees may include non-board advisors, volunteers, or community representatives. Committee chairs shall be appointed by the Board.")

section("Section 7.3 — Committee Authority")
body("Committees shall have only the authority delegated to them by the Board of Directors. No committee shall take action that binds the Corporation without Board approval.")

section("Section 7.4 — Advisory Board")
body("The Board may establish an Advisory Board consisting of non-voting advisors who provide expertise in AI/technology, community development, education, nonprofit management, and other relevant areas. Advisory board members shall be appointed by the Board of Directors.")

# ARTICLE VIII
article("ARTICLE VIII — FINANCIAL MANAGEMENT")

section("Section 8.1 — Fiscal Year")
body("The fiscal year of the Corporation shall begin on January 1 and end on December 31 of each year.")

section("Section 8.2 — Contracts and Expenditures")
body("All contracts and financial commitments above $1,000 shall require approval by the President. Expenditures above $5,000 shall require approval by the Board of Directors.")

section("Section 8.3 — Checks and Deposits")
body("All funds of the Corporation shall be deposited in accounts maintained in the name of the Corporation. Checks, drafts, or orders for payment shall be signed by the Treasurer and/or President, or such other persons as the Board may authorize.")

section("Section 8.4 — Loans")
body("No loans shall be made by the Corporation to any director or officer.")

section("Section 8.5 — Audit")
body("The Board of Directors shall cause the financial statements of the Corporation to be reviewed annually. When annual revenues exceed $250,000, the Board shall engage an independent auditor for a full audit.")

section("Section 8.6 — Grants and Restricted Funds")
body("All grant funds and restricted contributions shall be tracked separately, used only for their stated purpose, and reported in accordance with grantor requirements and applicable accounting standards.")

# ARTICLE IX
article("ARTICLE IX — INDEMNIFICATION AND LIABILITY")

section("Section 9.1 — Indemnification")
body("The Corporation shall indemnify its directors and officers to the fullest extent permitted by the South Carolina Nonprofit Corporation Act against any claims, liabilities, judgments, fines, settlements, and reasonable expenses (including attorneys' fees) arising from their service to the Corporation, provided they acted in good faith and in a manner they reasonably believed to be in the best interests of the Corporation.")

section("Section 9.2 — Insurance")
body("The Board of Directors may authorize the purchase of directors and officers (D&O) liability insurance at its discretion.")

# ARTICLE X
article("ARTICLE X — RECORDS AND REPORTS")

section("Section 10.1 — Records")
body("The Corporation shall maintain the following records at its principal office or other designated location: Articles of Incorporation and all amendments; these Bylaws and all amendments; minutes of all Board and committee meetings; financial statements and tax filings (Form 990); Conflict of Interest Policy and annual disclosures; contracts, grants, and legal documents.")

section("Section 10.2 — Public Inspection")
body("The Corporation shall make its Form 1023 (or 1023-EZ), Form 990, and determination letter available for public inspection upon request, as required by federal law.")

# ARTICLE XI
article("ARTICLE XI — CONFLICT OF INTEREST")

section("Section 11.1 — Policy")
body("The Corporation shall adopt and maintain a written Conflict of Interest Policy. All directors and officers shall review and sign the policy annually and shall disclose any potential conflicts as they arise.")

section("Section 11.2 — Recusal")
body("Any director or officer with a personal financial interest in a matter before the Board shall disclose that interest and recuse themselves from deliberation and voting on that matter.")

# ARTICLE XII
article("ARTICLE XII — DISSOLUTION")

section("Section 12.1 — Dissolution")
body("The Corporation may be dissolved upon the affirmative vote of two-thirds (2/3) of the Board of Directors, provided such dissolution complies with applicable state and federal law.")

section("Section 12.2 — Distribution of Assets")
body("Upon dissolution, after paying or providing for all liabilities, the remaining assets shall be distributed to one or more organizations that are organized and operated exclusively for charitable or educational purposes and qualify as tax-exempt under Section 501(c)(3) of the Internal Revenue Code, as determined by the Board of Directors. Under no circumstances shall assets be distributed to any private individual.")

# ARTICLE XIII
article("ARTICLE XIII — AMENDMENTS")

section("Section 13.1 — Amendment Procedure")
body("These Bylaws may be amended, repealed, or replaced by a two-thirds (2/3) vote of the Board of Directors at any regular or special meeting, provided that written notice of the proposed amendment has been given to all directors at least ten (10) days before the meeting.")

# ARTICLE XIV
article("ARTICLE XIV — PARLIAMENTARY AUTHORITY")

section("Section 14.1 — Rules of Order")
body("The rules contained in the current edition of Robert's Rules of Order Newly Revised shall govern the Corporation in all cases in which they are applicable and in which they are not inconsistent with these Bylaws and any special rules of order the Board may adopt.")

doc.add_page_break()

# CERTIFICATION
article("CERTIFICATION")
body("We, the undersigned, being the initial directors of The SCAiL Initiative, Inc., do hereby certify that the foregoing Bylaws were adopted as the Bylaws of the Corporation.")
doc.add_paragraph()

sig_line("Joshua German", "Founder & President")
sig_line("Janel Moore", "Board Member, Treasurer")
sig_line("Kinsey Meggett, PhD", "Board Member, Secretary")
sig_line("Tre Jenkins", "Board Member, Technology & Data")
sig_line("Darren Burton", "Board Member, Education Policy & Operations")

doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("The SCAiL Initiative, Inc. — South Carolina AI Literacy")
r.font.size = Pt(9)
r.font.color.rgb = GRAY
p2 = doc.add_paragraph()
p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r2 = p2.add_run("SC Filing ID: 260320-1607524")
r2.font.size = Pt(9)
r2.font.color.rgb = GRAY

doc.save('/home/user/Non-Profit/docs/SCAiL-Bylaws.docx')
print("Done: docs/SCAiL-Bylaws.docx")
